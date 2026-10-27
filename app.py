import io
import os
import re
from typing import Any

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document

MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MAX_FILE_SIZE_MB = 5

st.set_page_config(page_title="Resume ATS Checker", page_icon="📄", layout="wide")


def extract_resume_text(uploaded_file) -> str:
    """Extract text from a PDF, DOCX, or TXT upload."""
    file_bytes = uploaded_file.getvalue()
    suffix = uploaded_file.name.lower().rsplit(".", 1)[-1]

    if suffix == "pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == "docx":
        document = Document(io.BytesIO(file_bytes))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif suffix == "txt":
        text = file_bytes.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError(
            "No readable text was found. If this is a scanned PDF, convert it to a text-based PDF or DOCX and try again."
        )
    return text


def estimate_basic_checks(resume_text: str) -> dict[str, Any]:
    """Simple transparent checks; this is not a real employer ATS score."""
    text = resume_text.lower()
    words = re.findall(r"\b[\w+#.-]+\b", text)
    section_patterns = {
        "Contact information": r"email|e-mail|phone|mobile|linkedin|github",
        "Experience": r"experience|employment|work history|internship",
        "Education": r"education|university|college|degree|bachelor|master",
        "Skills": r"skills|technical skills|competencies|technologies",
        "Projects": r"projects|personal project|academic project",
    }
    sections = {name: bool(re.search(pattern, text)) for name, pattern in section_patterns.items()}
    checks = {
        "Resume has readable text": len(words) >= 50,
        "Contact details appear present": sections["Contact information"],
        "Education section appears present": sections["Education"],
        "Skills section appears present": sections["Skills"],
        "Experience or project section appears present": sections["Experience"] or sections["Projects"],
        "Resume has a reasonable text length": 150 <= len(words) <= 1200,
    }
    return {"word_count": len(words), "sections": sections, "checks": checks}


def make_review_prompt(resume_text: str, job_description: str) -> str:
    job_context = job_description.strip() or "Not provided. Review the resume generally, without assuming a specific job."
    return f"""
You are a careful resume reviewer. Review the resume text below for ATS-readability and job relevance.

Important rules:
- Give an estimated ATS-readiness score from 0 to 100. It is an educational estimate, not a score from any real employer's ATS.
- Do not invent qualifications, experience, metrics, degrees, certifications, or skills.
- Distinguish clearly between issues visible in the supplied text and things you cannot verify.
- Be constructive, specific, and fair to students and entry-level candidates.
- If a job description is supplied, compare the resume with it and identify missing or weakly evidenced keywords. Never recommend keyword stuffing.
- Do not penalize a person for personal characteristics or infer protected attributes.
- Do not repeat sensitive contact details unnecessarily.

Return Markdown with these exact sections:
## Estimated ATS-readiness score
Give one integer score out of 100 and briefly explain that it is an estimate.
## Score breakdown
Give 4-6 criteria with a score and short reason. Use a total consistent with your overall score.
## Strengths
List the strongest points in the resume.
## Improvements to make first
Give prioritized, actionable changes.
## Keywords and job match
Compare with the job description if provided. Otherwise explain that job-specific matching cannot be assessed.
## Example rewrites
Provide up to 3 before/after examples using only facts already present in the resume. If details are missing, use clearly marked placeholders rather than inventing facts.
## Final checklist
